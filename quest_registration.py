"""Eventfrog import: allowlisted fields only; never evaluate spreadsheet formulas."""
from io import BytesIO
import re
from html import escape
import base64
import qrcode

MENTORING_ANNOTATIONS = {"rosie", "svial-mentoring", "svial mentoring", "rosie vom svial-mentoring", "mentor", "mentors", "mentorin", "mentor:in", "mentoren"}


def read_eventfrog(content):
    from openpyxl import load_workbook
    if len(content) > 10_000_000:
        raise ValueError("Verwende eine Excel-Datei mit weniger als 10 MB.")
    book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        for sheet in book:
            header = None
            result = []
            for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if number > 20000:
                    raise ValueError("Die Datei enthält mehr als 20’000 Zeilen.")
                values = [str(v).strip() if v is not None else "" for v in cells]
                if header is None:
                    labels = [v.casefold() for v in values]
                    if all(k in labels for k in ("vorname", "nachname", "e-mail")):
                        header = {k:labels.index(k) for k in ("vorname", "nachname", "e-mail", "ticket-id", "id", "annotation", "affiliation", "institution", "zugehörigkeit", "namensschild leer") if k in labels}
                    continue
                def get(key):
                    return values[header[key]] if key in header and header[key] < len(values) else ""
                if not any(get(k) for k in ("vorname", "nachname", "e-mail", "ticket-id", "id")):
                    continue
                result.append({"name":" ".join(filter(None,[get("vorname"),get("nachname")])), "email":get("e-mail"), "first_name":get("vorname"), "last_name":get("nachname"),
                               **({"affiliation":get(next(k for k in ("affiliation","institution","zugehörigkeit") if k in header))} if any(k in header for k in ("affiliation","institution","zugehörigkeit")) else {}),
                               "blank_badge":get("namensschild leer").casefold() in {"ja","true","1","x"},
                               "source_id":("ticket:"+get("ticket-id")) if get("ticket-id") else (("id:"+get("id")) if get("id") else ""),
                               **({"annotation":get("annotation")} if "annotation" in header else {}), "row":number, "invalid_name":not get("vorname") or not get("nachname")})
            if header is not None:
                return result
        raise ValueError("Keine Kopfzeile mit Vorname, Nachname und E-Mail gefunden.")
    finally:
        book.close()


def plan_import(rows, registrations):
    plan, seen_sources, seen_people = [], set(), set()
    known = dict(registrations)
    for index, source in enumerate(rows, 1):
        row = {k:str(source.get(k, "")).strip() for k in ("name", "email", "source_id", "annotation", "affiliation", "first_name", "last_name")}
        if "first_name" not in source:
            row["first_name"], _, row["last_name"] = row["name"].partition(" ")
        row.update(row=source.get("row",index), action="New", reason="", person=None)
        name, email, sid = row["name"], row["email"].casefold(), row["source_id"]
        identity = (name.casefold(), email)
        issue = ""
        if not name or source.get("invalid_name") or any(x.startswith("=") for x in (name,email,sid)):
            issue = "Vor- und Nachname sind Pflicht; Formeln werden nicht unterstützt."
        elif not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            issue = "Eine gültige E-Mail-Adresse ist erforderlich."
        elif (sid and sid in seen_sources) or (not sid and identity in seen_people):
            issue = "Doppelte Zeile in dieser Datei."
        matches = [p for p,r in registrations.items() if sid and r.get("source_id") == sid]
        if not sid:
            matches = [p for p,r in registrations.items() if (r["name"].casefold(), r["email"].casefold()) == identity]
            if not matches and any(r["email"].casefold()==email for r in registrations.values()):
                issue = "Diese E-Mail existiert mit einem anderen Namen. Prüfe die Person oder ergänze die Ticket-ID."
        if len(matches)>1: issue = "Mehrere Treffer in der Datenbank."
        previous = registrations[matches[0]] if matches else {}
        duplicates = [r for r in known.values() if (r.get("source_name",r["name"]).casefold(), r.get("source_email",r["email"]).casefold()) == identity]
        ordinal = len(duplicates)+1
        pending = bool(source.get("blank_badge")) or (not matches and bool(sid) and ordinal>1)
        row.update(source_name=name, source_email=row["email"], provisional_name=name+(" "+str(ordinal) if ordinal>1 else ""), blank_badge=pending, identity_pending=pending, identity_corrected=False)
        if matches:
            row.update(person=matches[0], action="Update")
            for key in ("affiliation", "annotation"):
                if key not in source: row[key]=previous.get(key, "")
            for key in ("source_name", "source_email", "provisional_name", "blank_badge", "identity_pending", "identity_corrected"):
                if key in previous: row[key]=previous[key]
            if previous.get("identity_corrected"):
                for key in ("name","email","first_name","last_name","affiliation"):
                    row[key]=previous.get(key,row[key])
            if not sid: row["source_id"]=previous.get("source_id", "")
        if source.get("blank_badge") and not row["identity_corrected"]:
            row.update(blank_badge=True, identity_pending=True)
        if row["identity_pending"]:
            row["name"] = row["provisional_name"]
        row["mapping"] = annotation_mapping(row["annotation"])["label"]
        if row["blank_badge"]: row["reason"]="Leeres Namensschild · tatsächliche Person unter Badge korrigieren erfassen."
        if issue: row.update(action="Review",reason=issue)
        plan.append(row)
        if row["action"] == "New": known["new:"+str(index)]=row
        if sid: seen_sources.add(sid)
        seen_people.add(identity)
    return plan


def print_documents(roster, people, base_url):
    fronts, slips = [], []
    for person in people:
        r = roster[person]
        from quest_badges import qr_image
        uri = "data:image/png;base64,"+base64.b64encode(qr_image(base_url.rstrip("/")+"/?badge="+person)).decode("ascii")
        fronts.append(f'<section class="badge"><header><h1>{escape("" if r.get("blank_badge") else r["name"])}</h1></header><p class="annotation">{escape(r.get("affiliation", ""))}</p><img class="qr" alt="Persönlicher öffentlicher QR-Code" src="{uri}"><p class="badge-id">{escape(r["id"])}</p><small>AGRO-FOOD MESSE 2026</small></section>')
        slips.append(f'<section><h2>PRIVAT · Zugangscode</h2><h1>{escape(r["name"])}</h1><p>{escape(base_url)}</p><strong>{escape(r["code"])}</strong><p>Separat im Badgehalter aufbewahren. Nicht öffentlich zeigen.</p></section>')
    style='<meta charset="utf-8"><style>@page{size:A4;margin:12mm}body{font:15px Arial;display:flex;flex-wrap:wrap;gap:5mm}section{box-sizing:border-box;width:86mm;height:110mm;border:1px solid #ccc;text-align:center;padding:6mm;break-inside:avoid}header{display:flex;align-items:start;justify-content:space-between;gap:3mm;min-height:18mm}h1{font-size:22px;color:#009641;margin:0;text-align:left;overflow-wrap:anywhere}.logo{width:16mm;height:auto}.qr{width:48mm;height:48mm}.annotation{text-align:left;height:10mm;margin:2mm 0;font-size:14px;overflow-wrap:anywhere}.badge-id{font-size:18px;letter-spacing:1px;margin:2mm}strong{overflow-wrap:anywhere}</style>'
    return tuple('<!doctype html><html lang="de">'+style+'<body>'+''.join(parts)+'</body></html>' for parts in (fronts,slips))


def annotation_mapping(value):
    """Exact, reviewable aliases; never infer affiliation from a loose substring."""
    from quest_core import ORGANISATIONS, CLUSTER_LABELS
    value = value.strip().casefold()
    aliases = {"coop":"in-3p9d", "lidl":"re-lidl", "lidl schweiz":"re-lidl", "svial":"sv-5w8j"}
    aliases.update({r[1].casefold():token for token,r in ORGANISATIONS.items()})
    if value in MENTORING_ANNOTATIONS:
        return {"company":"sv-5w8j", "label":"SVIAL · Mentoring-Quest"}
    if value in aliases:
        token=aliases[value]
        return {"company":token,"label":ORGANISATIONS[token][1]+" · "+CLUSTER_LABELS[ORGANISATIONS[token][2]]}
    return {"company":None,"label":"Freie Annotation · keine Quest-Zuordnung" if value else "Teilnehmer:in ohne Firmenzuordnung"}


def batch_archive(roster, people, base_url, mirror_backs=True):
    from zipfile import ZipFile, ZIP_DEFLATED
    badges, slips = print_documents(roster, people, base_url)
    stream=BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        from quest_badges import badge_docx
        archive.writestr("00-PRIVATE-template-badges.docx", badge_docx(roster, people, base_url, mirror_backs))
        archive.writestr("01-public-badges.html", badges)
        archive.writestr("02-PRIVATE-login-slips.html", slips)
        archive.writestr("READ-ME.txt", "Enthält private Zugangscodes. Vertraulich aufbewahren. Word-Vorlage: ungerade Seiten vorne, gerade Seiten hinten, 10 Badges pro Blatt. A4 bei 100 % drucken, bei gespiegelten Rückseiten Duplex an der langen Kante. Zuerst auf Normalpapier testen. HTML-Dateien bieten alternativ separate Zugangszettel.")
    return stream.getvalue()


def mock_eventfrog_xlsx(count=10):
    """Downloadable fictional fixture for exercising the real Excel import."""
    from openpyxl import Workbook
    names=[('Lea','Meier',''),('Alex','Keller','Coop'),('Noah','Frei','Lidl'),('Mia','Baumann','SVIAL'),('Jonas','Weber',''),('Sara','Rossi','Mentorin'),('Luca','Huber','mooh'),('Anna','Müller','Emmi'),('Nina','Graf',''),('Tim','Steiner','SVIAL'),('Eva','Kunz','Strickhof'),('Max','Berger','')]
    book=Workbook();sheet=book.active;sheet.title='Fictional Eventfrog sample'
    sheet.append([f'AFJD fictional badge rehearsal — {count} participants'])
    sheet.append(['Nur Testdaten. Beim Import erzeugt die App gültige IDs und Zugangscodes.'])
    sheet.append([])
    sheet.append(['Ticket-ID','Vorname','Nachname','E-Mail','Affiliation','Annotation'])
    annotations=['','Coop','Lidl','SVIAL','','Mentor','mooh Genossenschaft','Emmi Schweiz AG','','SVIAL','Strickhof','']
    for i,(first,last,affiliation) in enumerate(names[:count]):
        sheet.append([f'BADGE-MOCK-{i+1:02}',first,last,f'badge-mock-{i+1:02}@example.test',affiliation,annotations[i]])
    for col,width in [('A',23),('B',18),('C',18),('D',34),('E',22),('F',29)]:sheet.column_dimensions[col].width=width
    stream=BytesIO();book.save(stream);return stream.getvalue()
